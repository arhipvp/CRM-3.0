"""User-scoped file access for MCP; reading never creates CRM or Drive records."""

from __future__ import annotations

import base64
import hashlib
import importlib
import io
import ipaddress
import mimetypes
import socket
import zipfile
from pathlib import PurePath
from types import SimpleNamespace
from urllib.parse import urljoin, urlsplit

import certifi
import urllib3
from apps.common import drive
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core import signing
from django.core.files.base import ContentFile
from django.http import FileResponse, Http404, QueryDict
from django.shortcuts import get_object_or_404
from django.utils.text import get_valid_filename
from rest_framework.exceptions import (
    AuthenticationFailed,
    NotFound,
    PermissionDenied,
    ValidationError,
)

from .contracts import Principal

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_CATALOG_FILES = 5000
MAX_IMAGE_PIXELS = 25_000_000
MAX_TEXT_CHARS = 50_000
DOWNLOAD_SALT = "crm.mcp.file-download.v1"
OWNER_TYPES = (
    "deal",
    "client",
    "policy",
    "statement",
    "note",
    "document",
    "ai_document",
)


def _check_access(principal, *, write=False, ai=False):
    from .auth import require_access

    scopes = {"crm.files", "crm.write" if write else "crm.read"}
    if ai:
        scopes.add("crm.ai")
    return require_access(principal, scopes)


def _view(view_class, user):
    view = view_class()
    view.action = "retrieve"
    view.request = SimpleNamespace(
        user=user,
        query_params=QueryDict("include_closed=true"),
        method="GET",
    )
    view.kwargs = {}
    return view


def _owner(owner_type, owner_id, principal, *, write=False):
    """Use each section's actual user queryset and mutation rule."""
    principal = _check_access(principal, write=write, ai=owner_type == "ai_document")
    user = principal.user
    if owner_type == "deal":
        from apps.deals.models import Deal
        from apps.deals.permissions import (
            build_deal_visibility_q,
            can_modify_deal,
            is_admin_user,
        )

        queryset = Deal.objects.select_related("client")
        if not is_admin_user(user):
            queryset = queryset.filter(build_deal_visibility_q(user)).distinct()
        obj = get_object_or_404(queryset, pk=owner_id)
        if write and not can_modify_deal(user, obj):
            raise PermissionDenied("You cannot modify this deal.")
        return obj
    if owner_type == "ai_document":
        from apps.ai_assistant.models import AiDocument
        from apps.ai_assistant.views import _can_manage_library

        obj = get_object_or_404(AiDocument, pk=owner_id)
        if write and not _can_manage_library(user):
            raise PermissionDenied("You cannot manage the shared library.")
        return obj

    view_classes = {
        "client": ("apps.clients.views", "ClientViewSet"),
        "document": ("apps.documents.views", "DocumentViewSet"),
        "note": ("apps.notes.views", "NoteViewSet"),
        "policy": ("apps.policies.views", "PolicyViewSet"),
        "statement": ("apps.finances.views", "StatementViewSet"),
    }
    if owner_type not in view_classes:
        raise ValidationError("Unsupported file owner type.")
    module, name = view_classes[owner_type]
    view = _view(getattr(importlib.import_module(module), name), user)
    obj = get_object_or_404(view.get_queryset().distinct(), pk=owner_id)
    if write and not view._can_modify(user, obj):
        raise PermissionDenied("You cannot modify this file owner.")
    return obj


def _drive_file_id(raw_id):
    return "d_" + base64.urlsafe_b64encode(str(raw_id).encode()).decode().rstrip("=")


def _raw_drive_id(file_id):
    if not file_id.startswith("d_"):
        raise NotFound("Drive file not found.")
    try:
        value = base64.urlsafe_b64decode(file_id[2:] + "=" * (-len(file_id[2:]) % 4))
        raw_id = value.decode("utf-8")
    except (ValueError, UnicodeDecodeError) as exc:
        raise NotFound("Drive file not found.") from exc
    if _drive_file_id(raw_id) != file_id:
        raise NotFound("Drive file not found.")
    return raw_id


def _tree(folder_id):
    if not folder_id:
        return {}
    result = {}
    pending, visited = [folder_id], set()
    while pending:
        folder = pending.pop()
        if folder in visited:
            continue
        visited.add(folder)
        for item in drive.list_drive_folder_contents(folder):
            result[item["id"]] = item
            if len(result) > MAX_CATALOG_FILES:
                raise ValidationError("File catalog exceeds 5000 entries.")
            if item["is_folder"]:
                pending.append(item["id"])
    return result


def _drive_entry(item):
    return {
        "id": _drive_file_id(item["id"]),
        "name": item["name"],
        "mime_type": item["mime_type"] or "application/octet-stream",
        "size": item["size"],
        "is_folder": bool(item["is_folder"]),
        "modified_at": item.get("modified_at"),
        "parent_id": (
            _drive_file_id(item["parent_id"]) if item.get("parent_id") else None
        ),
    }


def _local_entry(obj, *, ai=False):
    filename = obj.filename if ai else obj.title or PurePath(obj.file.name).name
    return {
        "id": ("a_" if ai else "l_") + str(obj.pk),
        "name": filename,
        "mime_type": ("" if ai else obj.mime_type)
        or mimetypes.guess_type(filename)[0]
        or "application/octet-stream",
        "size": obj.file.size,
        "is_folder": False,
    }


def _catalog(owner_type, obj):
    if owner_type in {"document", "ai_document"}:
        return [_local_entry(obj, ai=owner_type == "ai_document")] if obj.file else []
    if owner_type == "deal":
        from apps.codex_reader.views import _catalog as deal_catalog

        return [dict(item, is_folder=False) for item in deal_catalog(obj)]
    if owner_type == "note":
        if not obj.deal_id:
            return []
        tree = _tree(obj.deal.drive_folder_id)
        allowed = {
            str(item.get("id")) for item in obj.attachments if isinstance(item, dict)
        }
        return [
            _drive_entry(item)
            for key, item in tree.items()
            if key in allowed and not item["is_folder"]
        ]
    return [_drive_entry(item) for item in _tree(obj.drive_folder_id).values()]


def list_files(arguments, principal):
    owner_type, owner_id = arguments["owner_type"], arguments["owner_id"]
    obj = _owner(owner_type, owner_id, principal)
    items = _catalog(owner_type, obj)
    page, size = arguments.get("page", 1), arguments.get("page_size", 50)
    if not 1 <= page <= 100000 or not 1 <= size <= 100:
        raise ValidationError("Invalid pagination.")
    offset = (page - 1) * size
    return {
        "owner_type": owner_type,
        "owner_id": str(obj.pk),
        "count": len(items),
        "root_folder_id": (
            _drive_file_id(obj.drive_folder_id)
            if getattr(obj, "drive_folder_id", None)
            else None
        ),
        "next_page": page + 1 if offset + size < len(items) else None,
        "results": items[offset : offset + size],
    }


def _load_file(arguments, principal):
    owner_type, owner_id, file_id = (
        arguments["owner_type"],
        arguments["owner_id"],
        arguments["file_id"],
    )
    obj = _owner(owner_type, owner_id, principal)
    entry = next(
        (item for item in _catalog(owner_type, obj) if item["id"] == file_id), None
    )
    if entry is None or entry.get("is_folder"):
        raise NotFound("File does not belong to this owner.")
    if entry.get("size") is not None and int(entry["size"]) > MAX_FILE_BYTES:
        raise ValidationError("File exceeds the 25 MiB limit.")
    if file_id.startswith("d_"):
        content = _download_drive_bounded(_raw_drive_id(file_id))
    else:
        if owner_type not in {"document", "ai_document"}:
            from apps.documents.models import Document

            obj = get_object_or_404(Document, pk=file_id[2:], deal_id=owner_id)
        if obj.file.size > MAX_FILE_BYTES:
            raise ValidationError("File exceeds the 25 MiB limit.")
        with obj.file.open("rb") as stream:
            content = stream.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise ValidationError("File exceeds the 25 MiB limit.")
    return entry, content


def _download_drive_bounded(file_id):
    if drive._MediaIoBaseDownload is None:
        raise drive.DriveConfigurationError(
            "Drive download dependencies are unavailable."
        )

    def read(service):
        request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
        buffer = io.BytesIO()
        downloader = drive._MediaIoBaseDownload(buffer, request, chunksize=1024 * 1024)
        done = False
        while not done:
            _, done = downloader.next_chunk()
            if buffer.tell() > MAX_FILE_BYTES:
                raise ValidationError("File exceeds the 25 MiB limit.")
        return buffer.getvalue()

    return drive._run_with_drive_service("mcp_download_file", read)


def _validate_download_url(url):
    try:
        parsed = urlsplit(url)
        port = parsed.port
    except ValueError as exc:
        raise ValidationError("Invalid file download URL.") from exc
    host = (parsed.hostname or "").lower().rstrip(".")
    allowed = getattr(settings, "MCP_FILE_DOWNLOAD_HOSTS", ("oaiusercontent.com",))
    if (
        parsed.scheme != "https"
        or parsed.username
        or parsed.password
        or port not in {None, 443}
        or parsed.fragment
        or not any(host == value or host.endswith("." + value) for value in allowed)
    ):
        raise ValidationError("File URL must use an approved HTTPS download host.")
    try:
        addresses = list(
            dict.fromkeys(item[4][0] for item in socket.getaddrinfo(host, 443))
        )
    except OSError as exc:
        raise ValidationError("File download host is unavailable.") from exc
    if not addresses or any(
        not ipaddress.ip_address(value).is_global for value in addresses
    ):
        raise ValidationError(
            "Private, loopback and reserved download addresses are forbidden."
        )
    return parsed, host, addresses[0]


def download_chatgpt_file(file):
    """Pin the validated public address while retaining TLS hostname verification."""
    if (
        not isinstance(file, dict)
        or not file.get("file_id")
        or not file.get("download_url")
    ):
        raise ValidationError("A ChatGPT file_id and download_url are required.")
    url = file["download_url"]
    for _ in range(4):
        parsed, host, address = _validate_download_url(url)
        pool = urllib3.HTTPSConnectionPool(
            address,
            port=443,
            server_hostname=host,
            assert_hostname=host,
            cert_reqs="CERT_REQUIRED",
            ca_certs=certifi.where(),
            timeout=urllib3.Timeout(connect=5, read=20),
        )
        response = None
        try:
            response = pool.urlopen(
                "GET",
                parsed.path + ("?" + parsed.query if parsed.query else ""),
                headers={"Host": host},
                redirect=False,
                retries=False,
                preload_content=False,
            )
            if response.status in {301, 302, 303, 307, 308}:
                location = response.headers.get("Location")
                if not location:
                    raise ValidationError("Invalid file download redirect.")
                url = urljoin(url, location)
                continue
            if response.status != 200:
                raise ValidationError("ChatGPT file download failed.")
            length = response.headers.get("Content-Length")
            if length and int(length) > MAX_FILE_BYTES:
                raise ValidationError("File exceeds the 25 MiB limit.")
            data = bytearray()
            for chunk in response.stream(64 * 1024, decode_content=False):
                data.extend(chunk)
                if len(data) > MAX_FILE_BYTES:
                    raise ValidationError("File exceeds the 25 MiB limit.")
            return bytes(data)
        except (urllib3.exceptions.HTTPError, OSError, ValueError) as exc:
            raise ValidationError("ChatGPT file download failed.") from exc
        finally:
            if response is not None:
                response.close()
            pool.close()
    raise ValidationError("Too many file download redirects.")


def _filename(value):
    if not value or len(value) > 240 or "/" in value or "\\" in value:
        raise ValidationError(
            "Provide a file name without a path, up to 240 characters."
        )
    try:
        return get_valid_filename(value)
    except Exception as exc:
        raise ValidationError("Invalid file name.") from exc


def _audit(obj, principal, operation, details):
    from apps.users.models import AuditLog

    object_type = obj._meta.model_name
    AuditLog.objects.create(
        actor=principal.user,
        object_type="document" if object_type == "aidocument" else object_type,
        object_id=str(obj.pk),
        object_name=str(obj)[:255],
        action="update",
        description=f"ChatGPT: {operation}",
        new_value={"channel": "ChatGPT", **details},
    )


def upload_file(arguments, principal):
    owner_type, owner_id = arguments["owner_type"], arguments["owner_id"]
    obj = _owner(owner_type, owner_id, principal, write=True)
    file = arguments["file"]
    filename = _filename(file.get("file_name") or arguments.get("name") or "upload")
    mime_type = (
        file.get("mime_type")
        or mimetypes.guess_type(filename)[0]
        or "application/octet-stream"
    )
    destination = arguments.get("destination", "drive")
    if owner_type not in {"deal", "client", "policy", "statement", "note"}:
        raise ValidationError("This owner does not support file upload.")
    if destination == "local" and owner_type != "deal":
        raise ValidationError("Local documents must belong to a deal.")
    content = download_chatgpt_file(file)
    if destination == "local":
        from apps.documents.models import Document

        document = Document(
            deal=obj,
            owner=principal.user,
            title=filename,
            file_size=len(content),
            mime_type=mime_type,
            checksum=hashlib.sha256(content).hexdigest(),
        )
        document._audit_actor = principal.user
        document.file.save(filename, ContentFile(content), save=False)
        document.save()
        _audit(document, principal, "upload local file", {"file_name": filename})
        return _local_entry(document)
    if destination != "drive":
        raise ValidationError("Unsupported upload destination.")
    ensure = {
        "deal": drive.ensure_deal_folder,
        "client": drive.ensure_client_folder,
        "policy": drive.ensure_policy_folder,
        "statement": drive.ensure_statement_folder,
        "note": lambda note: drive.ensure_deal_folder(note.deal),
    }
    folder_id = ensure[owner_type](obj)
    if not folder_id:
        raise ValidationError("Drive folder is not configured.")
    uploaded = drive.upload_file_to_drive(
        folder_id, io.BytesIO(content), filename, mime_type
    )
    if owner_type == "note":
        obj.attachments = [*(obj.attachments or []), uploaded]
        obj._audit_actor = principal.user
        obj.save(update_fields=["attachments", "updated_at"])
    _audit(
        obj,
        principal,
        "upload Drive file",
        {"file_id": uploaded["id"], "file_name": filename},
    )
    return _drive_entry(uploaded)


def upload_library_file(arguments, principal):
    from apps.ai_assistant.models import AiDocument
    from apps.ai_assistant.views import _can_manage_library, document_payload

    principal = _check_access(principal, write=True, ai=True)
    if not _can_manage_library(principal.user):
        raise PermissionDenied("You cannot manage the shared library.")
    file = arguments["file"]
    filename = _filename(file.get("file_name") or arguments.get("name") or "upload")
    content = download_chatgpt_file(file)
    document = AiDocument(filename=filename)
    document.file.save(filename, ContentFile(content), save=False)
    document.save()
    _audit(
        document,
        principal,
        "upload AI library file; indexing queued",
        {"file_name": filename},
    )
    return document_payload(document)


def _replay_uploaded_file(arguments, principal, result):
    obj = _owner(arguments["owner_type"], arguments["owner_id"], principal, write=True)
    entries = _catalog(arguments["owner_type"], obj)
    if not any(entry["id"] == result.get("id") for entry in entries):
        raise NotFound("The uploaded file no longer belongs to this owner.")


def _replay_uploaded_library_file(arguments, principal, result):
    _owner("ai_document", result.get("id"), principal, write=True)


def list_library_files(arguments, principal):
    from apps.ai_assistant.models import AiDocument
    from apps.ai_assistant.views import document_payload

    _check_access(principal, ai=True)
    page, size = arguments.get("page", 1), arguments.get("page_size", 50)
    if page < 1 or not 1 <= size <= 100:
        raise ValidationError("Invalid pagination.")
    queryset = AiDocument.objects.all()
    query = arguments.get("query", "").strip()
    if query:
        queryset = queryset.filter(filename__icontains=query)
    start, count = (page - 1) * size, queryset.count()
    return {
        "count": count,
        "next_page": page + 1 if start + size < count else None,
        "results": [document_payload(item) for item in queryset[start : start + size]],
    }


def _ocr(image):
    import pytesseract

    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise ValidationError("Image exceeds the 25 megapixel OCR limit.")
    try:
        return pytesseract.image_to_string(image, lang="rus+eng", timeout=20)
    except (pytesseract.TesseractError, RuntimeError, OSError) as exc:
        raise ValidationError(
            "Local OCR is unavailable or exceeded its time limit."
        ) from exc


def _text_pages(content, filename, mime_type, start, count, use_ocr):
    """Read selected pages only; no CRM recognition or paid model is invoked."""
    import fitz
    from PIL import Image, ImageOps

    if content.startswith(b"%PDF") or mime_type == "application/pdf":
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                if document.needs_pass:
                    raise ValidationError("Password-protected PDF cannot be read.")
                total = document.page_count
                pages = []
                for index in range(start - 1, min(start - 1 + count, total)):
                    page = document[index]
                    text = page.get_text(sort=True)
                    ocr_used = not text.strip() and use_ocr
                    if ocr_used:
                        scale = min(2, 2000 / max(page.rect.width, page.rect.height))
                        pixmap = page.get_pixmap(
                            matrix=fitz.Matrix(scale, scale), alpha=False
                        )
                        text = _ocr(
                            Image.frombytes(
                                "RGB", (pixmap.width, pixmap.height), pixmap.samples
                            )
                        )
                    pages.append({"page": index + 1, "text": text, "ocr": ocr_used})
                return pages, total
        except (fitz.FileDataError, fitz.EmptyFileError) as exc:
            raise ValidationError("PDF cannot be decoded.") from exc
    if mime_type.startswith("image/"):
        try:
            with Image.open(io.BytesIO(content)) as original:
                if original.width * original.height > MAX_IMAGE_PIXELS:
                    raise ValidationError("Image exceeds the 25 megapixel limit.")
                text = _ocr(ImageOps.exif_transpose(original)) if use_ocr else ""
        except (Image.DecompressionBombError, OSError) as exc:
            raise ValidationError("Image cannot be decoded.") from exc
        return ([{"page": 1, "text": text, "ocr": use_ocr}] if start == 1 else []), 1
    suffix = PurePath(filename).suffix.lower()
    if suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 5000
                or sum(item.file_size for item in archive.infolist())
                > 100 * 1024 * 1024
            ):
                raise ValidationError("Office archive exceeds the extraction limit.")
        if suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            try:
                values = []
                for sheet in workbook.worksheets[:20]:
                    values.append(sheet.title)
                    for row in sheet.iter_rows(
                        max_row=1000, max_col=50, values_only=True
                    ):
                        values.append(
                            "\t".join(
                                str(value) if value is not None else "" for value in row
                            )
                        )
                text = "\n".join(values)
            finally:
                workbook.close()
    elif mime_type.startswith("text/") or suffix in {
        ".txt",
        ".csv",
        ".json",
        ".xml",
        ".md",
        ".eml",
    }:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1251")
    else:
        raise ValidationError(
            "This file type has no local text extractor; download the original."
        )
    return ([{"page": 1, "text": text, "ocr": False}] if start == 1 else []), 1


def read_file_text(arguments, principal):
    entry, content = _load_file(arguments, principal)
    start, count = arguments.get("start_page", 1), arguments.get("page_count", 3)
    offset = arguments.get("text_offset", 0)
    limit = arguments.get("max_chars", 20000)
    if (
        start < 1
        or not 1 <= count <= 5
        or offset < 0
        or not 100 <= limit <= MAX_TEXT_CHARS
    ):
        raise ValidationError("Invalid page or text range.")
    pages, total = _text_pages(
        content,
        entry["name"],
        entry["mime_type"],
        start,
        count,
        arguments.get("ocr", True),
    )
    remaining, result, next_page, next_offset = limit, [], None, 0
    for index, item in enumerate(pages):
        page_offset = offset if index == 0 else 0
        text = item["text"][page_offset:]
        result.append(dict(item, text=text[:remaining], text_offset=page_offset))
        if len(text) > remaining:
            next_page, next_offset = item["page"], page_offset + remaining
            break
        remaining -= len(text)
        if not remaining:
            next_page = item["page"] + 1 if item["page"] < total else None
            break
    else:
        next_page = (
            start + len(pages) if pages and start + len(pages) <= total else None
        )
    return {
        "source": {
            "owner_type": arguments["owner_type"],
            "owner_id": arguments["owner_id"],
            "file_id": entry["id"],
            "name": entry["name"],
        },
        "page_count": total,
        "pages": result,
        "next_page": next_page,
        "next_text_offset": next_offset,
    }


def read_file_page(arguments, principal):
    import fitz
    from mcp.types import CallToolResult, ImageContent
    from PIL import Image, ImageOps

    entry, content = _load_file(arguments, principal)
    page_number = arguments.get("page", 1)
    if content.startswith(b"%PDF") or entry["mime_type"] == "application/pdf":
        with fitz.open(stream=content, filetype="pdf") as document:
            if document.needs_pass or not 1 <= page_number <= document.page_count:
                raise ValidationError("PDF page is unavailable.")
            page = document[page_number - 1]
            scale = min(2, 2000 / max(page.rect.width, page.rect.height))
            image = page.get_pixmap(
                matrix=fitz.Matrix(scale, scale), alpha=False
            ).tobytes("png")
    elif entry["mime_type"].startswith("image/") and page_number == 1:
        with Image.open(io.BytesIO(content)) as original:
            if original.width * original.height > MAX_IMAGE_PIXELS:
                raise ValidationError("Image exceeds the 25 megapixel limit.")
            rendered = ImageOps.exif_transpose(original)
            rendered.thumbnail((2000, 2000))
            buffer = io.BytesIO()
            rendered.save(buffer, format="PNG")
            image = buffer.getvalue()
    else:
        raise ValidationError("Page preview is available for PDFs and images only.")
    return CallToolResult(
        content=[
            ImageContent(
                type="image",
                mimeType="image/png",
                data=base64.b64encode(image).decode(),
            )
        ]
    )


def get_file_download(arguments, principal):
    from mcp.types import CallToolResult, ResourceLink

    principal = _check_access(principal, ai=arguments["owner_type"] == "ai_document")
    obj = _owner(arguments["owner_type"], arguments["owner_id"], principal)
    entry = next(
        (
            item
            for item in _catalog(arguments["owner_type"], obj)
            if item["id"] == arguments["file_id"]
        ),
        None,
    )
    if entry is None or entry.get("is_folder"):
        raise NotFound("File does not belong to this owner.")
    if entry.get("size") is not None and int(entry["size"]) > MAX_FILE_BYTES:
        raise ValidationError("File exceeds the 25 MiB limit.")
    if principal.token_id is None:
        raise PermissionDenied(
            "A live OAuth token is required for file download links."
        )
    ticket = signing.dumps(
        {
            "user_id": principal.user.pk,
            "client_id": principal.client_id,
            "token_id": principal.token_id,
            "scopes": sorted(principal.scopes),
            "owner_type": arguments["owner_type"],
            "owner_id": arguments["owner_id"],
            "file_id": arguments["file_id"],
        },
        salt=DOWNLOAD_SALT,
        compress=True,
    )
    base_url = settings.MCP_PUBLIC_BASE_URL.rstrip("/")
    url = f"{base_url}/mcp-files/{ticket}/"
    return CallToolResult(
        content=[
            ResourceLink(
                type="resource_link",
                uri=url,
                name=entry["name"],
                mimeType=entry["mime_type"],
                description="Private file download link, valid for 10 minutes while access remains active.",
            )
        ]
    )


def file_download_view(request, ticket):
    """A signed capability grants only one file and keeps checking live OAuth access."""
    try:
        payload = signing.loads(ticket, salt=DOWNLOAD_SALT, max_age=600)
        principal = Principal(
            user=get_user_model().objects.get(pk=payload["user_id"], is_active=True),
            client_id=payload["client_id"],
            scopes=frozenset(payload["scopes"]),
            token_id=payload["token_id"],
        )
        principal = _check_access(principal, ai=payload["owner_type"] == "ai_document")
        entry, content = _load_file(payload, principal)
    except (
        signing.BadSignature,
        get_user_model().DoesNotExist,
        KeyError,
        ValueError,
        AuthenticationFailed,
        PermissionDenied,
        NotFound,
        ValidationError,
        drive.DriveError,
        Http404,
    ) as exc:
        raise Http404("File download is unavailable.") from exc
    safe_name = str(entry["name"]).replace("/", "_").replace("\\", "_")[:240]
    response = FileResponse(
        io.BytesIO(content),
        as_attachment=True,
        filename=_filename(safe_name or "download"),
        content_type=entry["mime_type"],
    )
    response["Cache-Control"] = "private, no-store"
    response["X-Content-Type-Options"] = "nosniff"
    response["Referrer-Policy"] = "no-referrer"
    return response


def list_file_folders(arguments, principal):
    obj = _owner(arguments["owner_type"], arguments["owner_id"], principal)
    if arguments["owner_type"] not in {"deal", "client", "policy", "statement"}:
        raise ValidationError("This owner has no Drive folders.")
    root = obj.drive_folder_id
    return {
        "root_folder_id": _drive_file_id(root) if root else None,
        "folders": [
            _drive_entry(item) for item in _tree(root).values() if item["is_folder"]
        ],
    }


def _mutation_plan(arguments, principal, operation):
    owner_type = arguments["owner_type"]
    allowed = {"deal", "statement"} if operation == "trash" else {"deal"}
    if owner_type not in allowed:
        raise ValidationError(
            "This owner does not support the selected file operation."
        )
    obj = _owner(owner_type, arguments["owner_id"], principal, write=True)
    tree = _tree(obj.drive_folder_id)
    raw_id = _raw_drive_id(arguments["file_id"])
    selected = tree.get(raw_id)
    if selected is None or selected["is_folder"]:
        raise NotFound("File does not belong to this owner's Drive folder.")
    changes = {"operation": operation, "file": _drive_entry(selected)}
    target = None
    if operation == "rename":
        changes["new_name"] = _filename(arguments["name"])
    if operation == "move":
        target = _raw_drive_id(arguments["target_folder_id"])
        if target != obj.drive_folder_id and (
            target not in tree or not tree[target]["is_folder"]
        ):
            raise PermissionDenied("Target folder must belong to the same deal.")
        changes["target_folder_id"] = _drive_file_id(target)
    return (
        obj,
        selected,
        target,
        {
            "fingerprint": {
                "owner_id": str(obj.pk),
                "owner_updated_at": obj.updated_at.isoformat(),
                "root_folder_id": obj.drive_folder_id,
                "file": selected,
                "target": tree.get(target) if target else None,
            },
            "changes": changes,
        },
    )


def _file_mutation(arguments, principal, operation):
    obj, selected, target, plan = _mutation_plan(arguments, principal, operation)
    if operation == "rename":
        result = drive.rename_drive_file(selected["id"], _filename(arguments["name"]))
    else:
        if operation == "trash":
            target = drive.ensure_trash_folder(obj.drive_folder_id)
        drive.move_drive_file_to_folder(selected["id"], target)
        result = dict(selected, parent_id=target)
    _audit(obj, principal, f"{operation} Drive file", plan["changes"])
    return {"file": _drive_entry(result), "operation": operation}


def get_file_specs():
    from .registry import ToolSpec, object_schema

    owner = {
        "owner_type": {"type": "string", "enum": list(OWNER_TYPES)},
        "owner_id": {"type": "string", "minLength": 1, "maxLength": 128},
    }
    selected = {
        **owner,
        "file_id": {"type": "string", "minLength": 1, "maxLength": 512},
    }
    required_owner = ["owner_type", "owner_id"]
    required_file = [*required_owner, "file_id"]
    file_schema = object_schema(
        {
            "download_url": {"type": "string", "format": "uri"},
            "file_id": {"type": "string", "minLength": 1, "maxLength": 255},
            "mime_type": {"type": "string", "maxLength": 120},
            "file_name": {"type": "string", "maxLength": 240},
        },
        ["download_url", "file_id"],
    )
    specs = [
        ToolSpec(
            "crm_list_files",
            "List files belonging to a CRM owner. Reading never creates Drive folders. Use the returned file IDs for text, page or original access.",
            object_schema(
                {
                    **owner,
                    "page": {"type": "integer", "minimum": 1},
                    "page_size": {"type": "integer", "minimum": 1, "maximum": 100},
                },
                required_owner,
            ),
            list_files,
            scopes=frozenset({"crm.read", "crm.files"}),
        ),
        ToolSpec(
            "crm_list_file_folders",
            "List existing Drive subfolders and the root folder ID for a client, deal, policy or finance statement; does not create folders.",
            object_schema(owner, required_owner),
            list_file_folders,
            scopes=frozenset({"crm.read", "crm.files"}),
        ),
        ToolSpec(
            "crm_read_file_text",
            "Read selected pages of an authorized PDF, image, text or Office file. Uses local OCR for scanned pages without changing CRM or invoking a paid model. Sources and continuation are returned.",
            object_schema(
                {
                    **selected,
                    "start_page": {"type": "integer", "minimum": 1},
                    "page_count": {"type": "integer", "minimum": 1, "maximum": 5},
                    "text_offset": {"type": "integer", "minimum": 0},
                    "max_chars": {
                        "type": "integer",
                        "minimum": 100,
                        "maximum": MAX_TEXT_CHARS,
                    },
                    "ocr": {"type": "boolean"},
                },
                required_file,
            ),
            read_file_text,
            scopes=frozenset({"crm.read", "crm.files"}),
        ),
        ToolSpec(
            "crm_read_file_page",
            "Render one authorized PDF page or image as an image for visual reading. No CRM recognition or state changes.",
            object_schema(
                {**selected, "page": {"type": "integer", "minimum": 1}}, required_file
            ),
            read_file_page,
            scopes=frozenset({"crm.read", "crm.files"}),
        ),
        ToolSpec(
            "crm_download_file",
            "Return a private download link for one authorized original file. The link expires in ten minutes and stops working when OAuth access is revoked.",
            object_schema(selected, required_file),
            get_file_download,
            scopes=frozenset({"crm.read", "crm.files"}),
        ),
        ToolSpec(
            "crm_upload_file",
            "Upload a file explicitly attached by the user to a CRM owner. May create its Drive folder. Local storage is supported only for deals. Does not overwrite existing files.",
            object_schema(
                {
                    **owner,
                    "file": file_schema,
                    "name": {"type": "string", "maxLength": 240},
                    "destination": {"type": "string", "enum": ["drive", "local"]},
                },
                [*required_owner, "file"],
            ),
            upload_file,
            scopes=frozenset({"crm.write", "crm.files"}),
            mutating=True,
            meta={"openai/fileParams": ["file"]},
            replay_guard=_replay_uploaded_file,
        ),
        ToolSpec(
            "crm_list_library_files",
            "List existing sources in the shared CRM AI document library. Does not call a paid model or index documents.",
            object_schema(
                {
                    "query": {"type": "string", "maxLength": 255},
                    "page": {"type": "integer", "minimum": 1},
                    "page_size": {"type": "integer", "minimum": 1, "maximum": 100},
                }
            ),
            list_library_files,
            scopes=frozenset({"crm.read", "crm.files", "crm.ai"}),
        ),
        ToolSpec(
            "crm_upload_library_file",
            "Upload an explicitly attached file to the shared AI library, subject to library manager rights. Queues the existing indexing process, which can use the configured paid provider.",
            object_schema(
                {"file": file_schema, "name": {"type": "string", "maxLength": 240}},
                ["file"],
            ),
            upload_library_file,
            scopes=frozenset({"crm.write", "crm.files", "crm.ai"}),
            mutating=True,
            meta={"openai/fileParams": ["file"]},
            replay_guard=_replay_uploaded_library_file,
        ),
    ]
    for operation, extra, required in (
        (
            "rename",
            {"name": {"type": "string", "minLength": 1, "maxLength": 240}},
            ["name"],
        ),
        (
            "move",
            {"target_folder_id": {"type": "string", "minLength": 1, "maxLength": 512}},
            ["target_folder_id"],
        ),
        ("trash", {}, []),
    ):
        specs.append(
            ToolSpec(
                f"crm_{operation}_drive_file",
                f"{operation.capitalize()} one existing Drive file belonging to a deal"
                + (" or finance statement" if operation == "trash" else "")
                + ". Requires a preview and explicit confirmation; folders and files outside the owner are rejected.",
                object_schema({**selected, **extra}, [*required_file, *required]),
                lambda arguments, principal, operation=operation: _file_mutation(
                    arguments, principal, operation
                ),
                scopes=frozenset({"crm.write", "crm.files", "crm.destructive"}),
                mutating=True,
                risky=True,
                previewer=lambda arguments, principal, operation=operation: _mutation_plan(
                    arguments, principal, operation
                )[
                    3
                ],
            )
        )
    return specs
