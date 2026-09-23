"""Extract indexed text and stable source locations from uploaded files."""

from __future__ import annotations

import csv
import email
import io
import tempfile
from dataclasses import dataclass
from email.message import Message
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    ".txt",
    ".md",
    ".csv",
    ".html",
    ".htm",
    ".eml",
    ".msg",
    ".jpg",
    ".jpeg",
    ".png",
    ".tiff",
    ".tif",
}


@dataclass(frozen=True)
class ExtractedPart:
    text: str
    location: dict[str, str | int]
    filename: str


def _ocr(data: bytes) -> str:
    import pytesseract
    from PIL import Image

    return pytesseract.image_to_string(
        Image.open(io.BytesIO(data)), lang="rus+eng"
    ).strip()


def _mail_text(message: Message) -> str:
    pieces = []
    for part in message.walk():
        if part.get_content_maintype() == "multipart" or part.get_filename():
            continue
        content = part.get_payload(decode=True)
        if content:
            pieces.append(
                content.decode(part.get_content_charset() or "utf-8", errors="replace")
            )
    return "\n".join(pieces)


def _attachment(name: str, content: bytes) -> list[ExtractedPart]:
    # Never trust an attachment name as a filesystem path.
    suffix = Path(name).suffix.lower()
    with tempfile.TemporaryDirectory() as directory:
        child = Path(directory) / f"attachment{suffix}"
        child.write_bytes(content)
        return [
            ExtractedPart(
                item.text,
                {"attachment_name": Path(name).name, **item.location},
                Path(name).name,
            )
            for item in extract(child)
        ]


def extract(path: Path, display_name: str | None = None) -> list[ExtractedPart]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Неподдерживаемый формат: {suffix}")
    name = display_name or path.name
    if suffix in {".txt", ".md"}:
        return [
            ExtractedPart(
                path.read_text(encoding="utf-8", errors="replace"),
                {"label": "текст"},
                name,
            )
        ]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as file:
            rows = [" | ".join(row) for row in csv.reader(file)]
        return [ExtractedPart("\n".join(rows), {"label": "CSV"}, name)]
    if suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        return [
            ExtractedPart(
                BeautifulSoup(path.read_bytes(), "html.parser").get_text(
                    "\n", strip=True
                ),
                {"label": "HTML"},
                name,
            )
        ]
    if suffix == ".pdf":
        import fitz

        result = []
        with fitz.open(path) as document:
            for number, page in enumerate(document, 1):
                content = page.get_text("text").strip()
                if not content:
                    content = _ocr(page.get_pixmap(dpi=200).pil_tobytes(format="PNG"))
                if content:
                    result.append(ExtractedPart(content, {"page": number}, name))
        return result
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        pieces = [item.text for item in document.paragraphs if item.text.strip()]
        for table in document.tables:
            pieces.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        return [ExtractedPart("\n".join(pieces), {"label": "DOCX"}, name)]
    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                ExtractedPart(
                    "\n".join(
                        " | ".join(
                            str(value) if value is not None else "" for value in row
                        )
                        for row in sheet.iter_rows(values_only=True)
                    ),
                    {"sheet": sheet.title},
                    name,
                )
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        return [
            ExtractedPart(
                "\n".join(
                    shape.text for shape in slide.shapes if hasattr(shape, "text")
                ),
                {"slide": number},
                name,
            )
            for number, slide in enumerate(presentation.slides, 1)
        ]
    if suffix in {".jpg", ".jpeg", ".png", ".tiff", ".tif"}:
        return [ExtractedPart(_ocr(path.read_bytes()), {"label": "OCR"}, name)]
    if suffix == ".eml":
        message = email.message_from_bytes(path.read_bytes())
        parts = [
            ExtractedPart(
                f"Тема: {message.get('Subject', '(без темы)')}\n{_mail_text(message)}",
                {"label": "письмо"},
                name,
            )
        ]
        for item in message.walk():
            attachment_name, content = item.get_filename(), item.get_payload(
                decode=True
            )
            if (
                attachment_name
                and content
                and Path(attachment_name).suffix.lower()
                in SUPPORTED_EXTENSIONS - {".eml", ".msg"}
            ):
                parts.extend(_attachment(attachment_name, content))
        return parts
    import extract_msg

    message = extract_msg.Message(str(path))
    parts = [
        ExtractedPart(
            f"Тема: {message.subject or '(без темы)'}\n{message.body or ''}",
            {"label": "письмо"},
            name,
        )
    ]
    for item in message.attachments:
        attachment_name = item.longFilename or item.shortFilename
        if (
            attachment_name
            and isinstance(item.data, bytes)
            and Path(attachment_name).suffix.lower()
            in SUPPORTED_EXTENSIONS - {".eml", ".msg"}
        ):
            parts.extend(_attachment(attachment_name, item.data))
    return parts
