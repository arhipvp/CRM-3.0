from __future__ import annotations

import csv
import email
import io
from dataclasses import dataclass
from email.message import Message
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    ".txt",
    ".md",
    ".csv",
    ".html",
    ".htm",
    ".eml",
    ".msg",
    ".jpg",
    ".jpeg",
    ".png",
    ".tiff",
    ".tif",
}


@dataclass(frozen=True)
class ExtractedPart:
    text: str
    location: str
    filename: str


def extract(path: Path) -> list[ExtractedPart]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [
            ExtractedPart(
                path.read_text(encoding="utf-8", errors="replace"), "текст", path.name
            )
        ]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as file:
            rows = [" | ".join(row) for row in csv.reader(file)]
        return [ExtractedPart("\n".join(rows), "CSV", path.name)]
    if suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        return [
            ExtractedPart(
                BeautifulSoup(path.read_bytes(), "html.parser").get_text(
                    "\n", strip=True
                ),
                "HTML",
                path.name,
            )
        ]
    if suffix == ".pdf":
        import fitz

        result: list[ExtractedPart] = []
        document = fitz.open(path)
        for number, page in enumerate(document, 1):
            text = page.get_text("text").strip()
            if not text:
                text = _ocr_image(page.get_pixmap(dpi=200).pil_tobytes(format="PNG"))
            if text:
                result.append(ExtractedPart(text, f"страница {number}", path.name))
        return result
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        return [ExtractedPart("\n".join(paragraphs), "DOCX", path.name)]
    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return [
            ExtractedPart(
                "\n".join(
                    " | ".join(str(v) if v is not None else "" for v in row)
                    for row in sheet.iter_rows(values_only=True)
                ),
                f"лист {sheet.title}",
                path.name,
            )
            for sheet in workbook.worksheets
        ]
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        return [
            ExtractedPart(
                "\n".join(
                    shape.text for shape in slide.shapes if hasattr(shape, "text")
                ),
                f"слайд {number}",
                path.name,
            )
            for number, slide in enumerate(presentation.slides, 1)
        ]
    if suffix in {".jpg", ".jpeg", ".png", ".tiff", ".tif"}:
        return [ExtractedPart(_ocr_image(path.read_bytes()), "OCR", path.name)]
    if suffix == ".eml":
        return _extract_eml(path)
    if suffix == ".msg":
        return _extract_msg(path)
    raise ValueError(f"Неподдерживаемый формат: {suffix}")


def _ocr_image(data: bytes) -> str:
    import pytesseract
    from PIL import Image

    return pytesseract.image_to_string(
        Image.open(io.BytesIO(data)), lang="rus+eng"
    ).strip()


def _extract_eml(path: Path) -> list[ExtractedPart]:
    message = email.message_from_bytes(path.read_bytes())
    subject = message.get("Subject", "(без темы)")
    text = _email_text(message)
    result = [ExtractedPart(f"Тема: {subject}\n{text}", "письмо", path.name)]
    for part in message.walk():
        name = part.get_filename()
        content = part.get_payload(decode=True)
        if (
            name
            and content
            and Path(name).suffix.lower() in SUPPORTED_EXTENSIONS - {".eml", ".msg"}
        ):
            result.extend(_extract_attachment(name, content, path.name))
    return result


def _email_text(message: Message) -> str:
    chunks: list[str] = []
    for part in message.walk():
        if part.get_content_maintype() == "multipart" or part.get_filename():
            continue
        data = part.get_payload(decode=True)
        if data:
            chunks.append(
                data.decode(part.get_content_charset() or "utf-8", errors="replace")
            )
    return "\n".join(chunks)


def _extract_attachment(name: str, content: bytes, parent: str) -> list[ExtractedPart]:
    import tempfile

    with tempfile.TemporaryDirectory() as directory:
        child = Path(directory) / name
        child.write_bytes(content)
        return [
            ExtractedPart(item.text, f"вложение {name}; {item.location}", parent)
            for item in extract(child)
        ]


def _extract_msg(path: Path) -> list[ExtractedPart]:
    import extract_msg

    message = extract_msg.Message(str(path))
    result = [
        ExtractedPart(
            f"Тема: {message.subject or '(без темы)'}\n{message.body or ''}",
            "письмо",
            path.name,
        )
    ]
    for attachment in message.attachments:
        name = attachment.longFilename or attachment.shortFilename
        data = attachment.data
        if (
            name
            and isinstance(data, bytes)
            and Path(name).suffix.lower() in SUPPORTED_EXTENSIONS - {".eml", ".msg"}
        ):
            result.extend(_extract_attachment(name, data, path.name))
    return result
